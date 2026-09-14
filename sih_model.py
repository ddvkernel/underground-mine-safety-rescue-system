import os
import sys
import math
import time
import datetime
import argparse
import numpy as np
import pandas as pd
import serial

# Ensure UTF-8 output on Windows consoles
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

from sklearn.ensemble import RandomForestClassifier
from sklearn.preprocessing import StandardScaler
from sklearn.metrics import accuracy_score, classification_report

import matplotlib
matplotlib.use("Agg")  # Non-interactive backend for generating graphics
import matplotlib.pyplot as plt

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, LineChart, Reference, Series
from openpyxl.drawing.image import Image as OpenPyXLImage


# ============================================================
# PATHS AND DIRECTORY MANAGEMENT
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
INSIGHTS_DIR = os.path.join(BASE_DIR, "insights")
os.makedirs(INSIGHTS_DIR, exist_ok=True)

# Primary Excel output path (as requested: mine_rescue.xlsx)
EXCEL_FILE = os.path.join(INSIGHTS_DIR, "mine_rescue.xlsx")

# Legacy fallback file if mine_rescue.xlsx does not exist yet
LEGACY_EXCEL_FILE = os.path.join(INSIGHTS_DIR, "mine_safety_data.xlsx")

# Locate training files in data/ or root directory
def resolve_data_path(filename):
    path_in_data = os.path.join(DATA_DIR, filename)
    if os.path.exists(path_in_data):
        return path_in_data
    path_in_root = os.path.join(BASE_DIR, filename)
    if os.path.exists(path_in_root):
        return path_in_root
    return path_in_data

CO_FILE = resolve_data_path("ethylene_CO.txt")
CH4_FILE = resolve_data_path("ethylene_methane.txt")


# ============================================================
# CONFIGURATION
# ============================================================

N_ROWS = 20000
WINDOW_SIZE = 20
FUTURE_HORIZON = 20
RANDOM_STATE = 42


# ============================================================
# RISK THRESHOLDS
# ============================================================

CO_WARNING = 50.0
CO_CRITICAL = 100.0

CH4_WARNING = 50.0
CH4_CRITICAL = 100.0

VIB_WARNING = 0.5
VIB_CRITICAL = 1.0


# ============================================================
# SERIAL CONFIGURATION
# ============================================================

SERIAL_PORT = os.environ.get("SERIAL_PORT", "COM5")       # Change if required
BAUD_RATE = int(os.environ.get("BAUD_RATE", 115200))
ML_INTERVAL_SECONDS = 5.0


# ============================================================
# LOAD TRAINING DATA
# ============================================================

print()
print("=" * 60)
print("LOADING TRAINING DATA")
print("=" * 60)

print("CO dataset  :", CO_FILE)

if os.path.exists(CO_FILE):
    co = pd.read_csv(
        CO_FILE,
        sep=r"\s+",
        skiprows=1,
        usecols=[0, 1],
        names=["Time", "CO_ppm"],
        nrows=N_ROWS,
        engine="python"
    )
    co = co.reset_index(drop=True)
    print("CO samples  :", len(co))
else:
    print(f"WARNING: CO dataset not found at {CO_FILE}. Generating baseline synthetic training data.")
    co = pd.DataFrame({
        "Time": np.arange(N_ROWS),
        "CO_ppm": np.clip(np.random.normal(loc=15.0, scale=8.0, size=N_ROWS), 0, 300)
    })

print()
print("CH4 dataset :", CH4_FILE)

if os.path.exists(CH4_FILE):
    ch4 = pd.read_csv(
        CH4_FILE,
        sep=r"\s+",
        skiprows=1,
        usecols=[0, 1],
        names=["Time", "CH4_ppm"],
        nrows=N_ROWS,
        engine="python"
    )
    ch4 = ch4.reset_index(drop=True)
    print("CH4 samples :", len(ch4))
else:
    print(f"WARNING: CH4 dataset not found at {CH4_FILE}. Generating baseline synthetic training data.")
    ch4 = pd.DataFrame({
        "Time": np.arange(N_ROWS),
        "CH4_ppm": np.clip(np.random.normal(loc=20.0, scale=10.0, size=N_ROWS), 0, 200)
    })


# ============================================================
# SYNTHETIC VIBRATION FOR TRAINING
# ============================================================

np.random.seed(RANDOM_STATE)

def generate_vibration(n):
    vibration = np.zeros(n)
    vibration[:] = np.random.normal(loc=0.15, scale=0.04, size=n)

    region1_start = int(n * 0.25)
    region1_end = int(n * 0.30)
    region2_start = int(n * 0.60)
    region2_end = int(n * 0.65)

    vibration[region1_start:region1_end] += 0.35
    vibration[region2_start:region2_end] += 0.75
    return np.clip(vibration, 0, None)

co_vibration = generate_vibration(len(co))
ch4_vibration = generate_vibration(len(ch4))


# ============================================================
# TEMPORAL FEATURES
# ============================================================

def create_temporal_features(gas_values, vibration_values, index):
    start = index - WINDOW_SIZE + 1
    end = index + 1

    gas_window = gas_values[start:end]
    vib_window = vibration_values[start:end]

    features = [
        gas_window[-1],
        vib_window[-1],
        np.mean(gas_window),
        np.mean(vib_window),
        np.max(gas_window),
        np.max(vib_window),
        np.min(gas_window),
        np.min(vib_window),
        gas_window[-1] - gas_window[0],
        vib_window[-1] - vib_window[0],
        np.std(gas_window),
        np.std(vib_window),
        np.mean(gas_window[-5:]),
        np.mean(vib_window[-5:])
    ]
    return features


# ============================================================
# FUTURE LABEL
# ============================================================

def create_future_label(gas_values, vibration_values, index, gas_type):
    future_start = index + 1
    future_end = index + 1 + FUTURE_HORIZON

    future_gas = gas_values[future_start:future_end]
    future_vibration = vibration_values[future_start:future_end]

    if len(future_gas) < FUTURE_HORIZON:
        return None

    if gas_type == "CO":
        critical_condition = (
            (future_gas >= CO_CRITICAL) |
            (future_vibration >= VIB_CRITICAL)
        )
        warning_condition = (
            (future_gas >= CO_WARNING) |
            (future_vibration >= VIB_WARNING)
        )
    else:
        critical_condition = (
            (future_gas >= CH4_CRITICAL) |
            (future_vibration >= VIB_CRITICAL)
        )
        warning_condition = (
            (future_gas >= CH4_WARNING) |
            (future_vibration >= VIB_WARNING)
        )

    if np.any(critical_condition):
        return "CRITICAL"
    elif np.any(warning_condition):
        return "WARNING"
    else:
        return "SAFE"


# ============================================================
# BUILD TRAINING DATA
# ============================================================

def build_training_data(gas_values, vibration_values, gas_type):
    X = []
    y = []

    first_index = WINDOW_SIZE - 1
    last_index = len(gas_values) - FUTURE_HORIZON - 1

    for i in range(first_index, last_index + 1):
        features = create_temporal_features(gas_values, vibration_values, i)
        label = create_future_label(gas_values, vibration_values, i, gas_type)

        if label is not None:
            X.append(features)
            y.append(label)

    return np.array(X), np.array(y)


# ============================================================
# CREATE TRAINING DATASETS
# ============================================================

print()
print("=" * 60)
print("CREATING CO TRAINING DATA")
print("=" * 60)
X_CO, y_CO = build_training_data(co["CO_ppm"].values, co_vibration, "CO")
print("CO training examples:", len(X_CO))

print()
print("=" * 60)
print("CREATING CH4 TRAINING DATA")
print("=" * 60)
X_CH4, y_CH4 = build_training_data(ch4["CH4_ppm"].values, ch4_vibration, "CH4")
print("CH4 training examples:", len(X_CH4))


# ============================================================
# TRAIN RANDOM FOREST MODEL
# ============================================================

def train_model(X, y, name):
    total = len(X)
    train_end = int(total * 0.70)
    validation_end = int(total * 0.85)

    X_train = X[:train_end]
    y_train = y[:train_end]

    X_validation = X[train_end:validation_end]
    y_validation = y[train_end:validation_end]

    X_test = X[validation_end:]
    y_test = y[validation_end:]

    scaler = StandardScaler()
    X_train_scaled = scaler.fit_transform(X_train)
    X_validation_scaled = scaler.transform(X_validation)
    X_test_scaled = scaler.transform(X_test)

    model = RandomForestClassifier(
        n_estimators=200,
        max_depth=12,
        class_weight="balanced",
        random_state=RANDOM_STATE
    )
    model.fit(X_train_scaled, y_train)

    validation_prediction = model.predict(X_validation_scaled)
    print()
    print("=" * 60)
    print(name, "VALIDATION")
    print("=" * 60)
    print("Accuracy:", accuracy_score(y_validation, validation_prediction))
    print(classification_report(y_validation, validation_prediction, zero_division=0))

    test_prediction = model.predict(X_test_scaled)
    print()
    print("=" * 60)
    print(name, "FINAL TEST")
    print("=" * 60)
    print("Accuracy:", accuracy_score(y_test, test_prediction))
    print(classification_report(y_test, test_prediction, zero_division=0))

    return model, scaler


# ============================================================
# TRAIN MODELS
# ============================================================

print()
print("=" * 60)
print("TRAINING CO MODEL")
print("=" * 60)
co_model, co_scaler = train_model(X_CO, y_CO, "CO")

print()
print("=" * 60)
print("TRAINING CH4 MODEL")
print("=" * 60)
ch4_model, ch4_scaler = train_model(X_CH4, y_CH4, "CH4")


# ============================================================
# REFERENCE HISTORY FOR LIVE STREAMING
# ============================================================

co_reference_gas = co["CO_ppm"].values[-WINDOW_SIZE:].copy()
ch4_reference_gas = ch4["CH4_ppm"].values[-WINDOW_SIZE:].copy()
co_reference_vibration = co_vibration[-WINDOW_SIZE:].copy()
ch4_reference_vibration = ch4_vibration[-WINDOW_SIZE:].copy()


# ============================================================
# INFERENCE HELPERS
# ============================================================

def create_test_features(reference_gas, reference_vibration, current_gas, current_vibration):
    gas_window = reference_gas.copy()
    vibration_window = reference_vibration.copy()

    gas_window[-1] = current_gas
    vibration_window[-1] = current_vibration

    features = [
        gas_window[-1],
        vibration_window[-1],
        np.mean(gas_window),
        np.mean(vibration_window),
        np.max(gas_window),
        np.max(vibration_window),
        np.min(gas_window),
        np.min(vibration_window),
        gas_window[-1] - gas_window[0],
        vibration_window[-1] - vibration_window[0],
        np.std(gas_window),
        np.std(vibration_window),
        np.mean(gas_window[-5:]),
        np.mean(vibration_window[-5:])
    ]
    return np.array(features).reshape(1, -1)


def get_probability(model, scaler, features):
    scaled = scaler.transform(features)
    probabilities = model.predict_proba(scaled)[0]
    return dict(zip(model.classes_, probabilities))


def fuse_predictions(co_prob, ch4_prob):
    safe = (co_prob.get("SAFE", 0) + ch4_prob.get("SAFE", 0)) / 2
    warning = (co_prob.get("WARNING", 0) + ch4_prob.get("WARNING", 0)) / 2
    critical = (co_prob.get("CRITICAL", 0) + ch4_prob.get("CRITICAL", 0)) / 2

    probabilities = {
        "SAFE": safe,
        "WARNING": warning,
        "CRITICAL": critical
    }
    prediction = max(probabilities, key=probabilities.get)
    return prediction, probabilities


def mq5_voltage_to_ch4(voltage):
    return float(np.clip((voltage / 3.3) * 150.0, 0, 150))


def mq7_voltage_to_co(voltage):
    return float(np.clip((voltage / 3.3) * 600.0, 0, 600))


# ============================================================
# COMPREHENSIVE MULTI-SHEET EXCEL INSIGHTS ENGINE
# ============================================================

def determine_hazard_category(co_val, ch4_val, vib_val, risk):
    reasons = []
    if co_val >= CO_CRITICAL:
        reasons.append("Critical Toxic CO")
    elif co_val >= CO_WARNING:
        reasons.append("Elevated CO")

    if ch4_val >= CH4_CRITICAL:
        reasons.append("Critical Explosive CH4")
    elif ch4_val >= CH4_WARNING:
        reasons.append("Elevated CH4")

    if vib_val >= VIB_CRITICAL:
        reasons.append("Critical Seismic Tremor")
    elif vib_val >= VIB_WARNING:
        reasons.append("Seismic Activity")

    if not reasons:
        return "Normal Air & Structure" if risk == "SAFE" else "ML Predictive Risk"
    return " & ".join(reasons)


def determine_safety_action(risk, hazard):
    if risk == "CRITICAL":
        if "CH4" in hazard or "CO" in hazard:
            return "IMMEDIATE EVACUATION: Engage Max Auxiliary Ventilation & Audio Sirens"
        elif "Seismic" in hazard:
            return "STRUCTURAL WARNING: Immediate Retreat to Reinforced Shelter / Shaft"
        return "CRITICAL EVACUATION: Immediate Sector Clearance & Search Party Dispatch"
    elif risk == "WARNING":
        if "CH4" in hazard:
            return "VENTILATION ALERT: Trigger Zone Exhauster & Inspect Gas Pocket"
        elif "CO" in hazard:
            return "AIR QUALITY ALERT: Deploy Respirators & Check Fire Dampers"
        return "ENHANCED MONITORING: Restrict Tunnel Entry & Prepare Standby Extraction"
    return "CONTINUOUS MONITORING: Environmental Parameters Normal"


# ============================================================
# VISUAL GRAPHING & ANALYTICS DASHBOARD GENERATION
# ============================================================

def generate_visual_dashboard(df_raw, output_path=None, sync_legacy=True):
    """
    Generates a publication-grade, multi-panel visual analytics dashboard PNG
    graphing risk breakdowns, gas progression, seismic vibration, zone comparisons,
    and worker signal tracking.
    """
    if df_raw is None or df_raw.empty:
        print("[VISUAL GRAPHS] Warning: Empty dataframe, skipping visual graph generation.")
        return None

    if output_path is None:
        output_path = os.path.join(INSIGHTS_DIR, "mine_rescue_dashboard.png")
    output_path = os.path.abspath(os.path.normpath(output_path))

    df = df_raw.copy()

    for col in ["CH4 (ppm)", "CO (ppm)", "Vibration", "RSSI (dBm)"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    if "Zone" in df.columns:
        df["Zone"] = pd.to_numeric(df["Zone"], errors="coerce").fillna(0).astype(int)

    # Styling setup
    plt.rcParams.update({
        "figure.facecolor": "#0B1220",
        "axes.facecolor": "#111827",
        "axes.edgecolor": "#334155",
        "axes.labelcolor": "#E2E8F0",
        "xtick.color": "#94A3B8",
        "ytick.color": "#94A3B8",
        "text.color": "#F8FAFC",
        "grid.color": "#1E293B",
        "grid.linestyle": "--",
        "grid.alpha": 0.6,
        "font.family": "sans-serif"
    })

    fig = plt.figure(figsize=(18, 12), dpi=200)
    fig.suptitle("UNDERGROUND MINE SAFETY & WORKER RESCUE ANALYTICS DASHBOARD",
                 fontsize=17, fontweight="bold", color="#38BDF8", y=0.98)

    gs = fig.add_gridspec(3, 2, hspace=0.35, wspace=0.25, top=0.92, bottom=0.06, left=0.07, right=0.96)

    # 1. DONUT CHART: Overall Mine Risk Distribution
    ax1 = fig.add_subplot(gs[0, 0])
    risk_counts = df["Risk Status"].value_counts()
    categories = ["SAFE", "WARNING", "CRITICAL"]
    counts = [risk_counts.get(c, 0) for c in categories]
    colors = ["#22C55E", "#F59E0B", "#EF4444"]

    plot_labels = [c for c, count in zip(categories, counts) if count > 0]
    plot_counts = [count for count in counts if count > 0]
    plot_colors = [col for c, count, col in zip(categories, counts, colors) if count > 0]

    if sum(plot_counts) > 0:
        wedges, texts, autotexts = ax1.pie(
            plot_counts,
            labels=plot_labels,
            autopct="%1.1f%%",
            startangle=140,
            colors=plot_colors,
            wedgeprops={"width": 0.45, "edgecolor": "#0B1220", "linewidth": 2},
            pctdistance=0.75,
            textprops={"fontsize": 10, "fontweight": "bold"}
        )
        for autotext in autotexts:
            autotext.set_color("#FFFFFF")
            autotext.set_fontsize(11)
        total_rec = len(df)
        safe_pct = (risk_counts.get("SAFE", 0) / total_rec * 100) if total_rec > 0 else 100
        ax1.text(0, 0, f"Safety Rating\n{safe_pct:.1f}%\n({total_rec} Cycles)",
                 ha="center", va="center", fontsize=11, fontweight="bold", color="#38BDF8")
    ax1.set_title("Operational Risk & Safety Distribution", fontsize=13, fontweight="bold", pad=10, color="#F1F5F9")

    # 2. TIME-SERIES: Toxic Gas Concentrations (CO & CH4)
    ax2 = fig.add_subplot(gs[0, 1])
    sample_indices = np.arange(len(df))
    if "CO (ppm)" in df.columns:
        ax2.plot(sample_indices, df["CO (ppm)"], label="Carbon Monoxide (CO)", color="#F87171", linewidth=1.5, alpha=0.9)
    if "CH4 (ppm)" in df.columns:
        ax2.plot(sample_indices, df["CH4 (ppm)"], label="Methane (CH4)", color="#FBBF24", linewidth=1.5, alpha=0.9)

    ax2.axhline(CO_WARNING, color="#F59E0B", linestyle=":", linewidth=1.2, label=f"Warning Limit ({CO_WARNING:.0f} ppm)")
    ax2.axhline(CO_CRITICAL, color="#EF4444", linestyle="--", linewidth=1.2, label=f"Critical Limit ({CO_CRITICAL:.0f} ppm)")
    ax2.set_title("Toxic & Explosive Gas Telemetry Progression", fontsize=13, fontweight="bold", pad=10, color="#F1F5F9")
    ax2.set_xlabel("Telemetry Cycle Index", fontsize=10)
    ax2.set_ylabel("Concentration (ppm)", fontsize=10)
    ax2.grid(True)
    ax2.legend(loc="upper right", fontsize=8, facecolor="#1E293B", edgecolor="#334155")

    # 3. SEISMIC VIBRATION: Vibration (g) Spike Profile
    ax3 = fig.add_subplot(gs[1, 0])
    if "Vibration" in df.columns:
        ax3.plot(sample_indices, df["Vibration"], color="#A78BFA", linewidth=1.2, label="Vibration (g)")
        ax3.fill_between(sample_indices, df["Vibration"], color="#8B5CF6", alpha=0.15)
    ax3.axhline(VIB_WARNING, color="#F59E0B", linestyle=":", linewidth=1.2, label=f"Warning ({VIB_WARNING:.1f} g)")
    ax3.axhline(VIB_CRITICAL, color="#EF4444", linestyle="--", linewidth=1.2, label=f"Critical ({VIB_CRITICAL:.1f} g)")
    ax3.set_title("Seismic Vibration & Structural Disturbance", fontsize=13, fontweight="bold", pad=10, color="#F1F5F9")
    ax3.set_xlabel("Telemetry Cycle Index", fontsize=10)
    ax3.set_ylabel("Acceleration Delta (g)", fontsize=10)
    ax3.grid(True)
    ax3.legend(loc="upper right", fontsize=8, facecolor="#1E293B", edgecolor="#334155")

    # 4. CROSS-ZONE COMPARISON: Zone 1 vs Zone 2 Metrics
    ax4 = fig.add_subplot(gs[1, 1])
    zone_ids = [z for z in sorted(df["Zone"].unique()) if z != 0]
    if zone_ids:
        x_pos = np.arange(len(zone_ids))
        width = 0.20

        co_means = [df[df["Zone"] == z]["CO (ppm)"].mean() for z in zone_ids]
        co_maxs = [df[df["Zone"] == z]["CO (ppm)"].max() for z in zone_ids]
        ch4_means = [df[df["Zone"] == z]["CH4 (ppm)"].mean() for z in zone_ids]
        ch4_maxs = [df[df["Zone"] == z]["CH4 (ppm)"].max() for z in zone_ids]

        ax4.bar(x_pos - 1.5*width, co_means, width, label="Mean CO", color="#F87171")
        ax4.bar(x_pos - 0.5*width, co_maxs, width, label="Peak CO", color="#DC2626")
        ax4.bar(x_pos + 0.5*width, ch4_means, width, label="Mean CH4", color="#FBBF24")
        ax4.bar(x_pos + 1.5*width, ch4_maxs, width, label="Peak CH4", color="#D97706")

        ax4.set_xticks(x_pos)
        ax4.set_xticklabels([f"Zone {z}" for z in zone_ids], fontsize=11, fontweight="bold")
        ax4.set_ylabel("Gas PPM", fontsize=10)
        ax4.set_title("Cross-Zone Environmental Gas Comparison", fontsize=13, fontweight="bold", pad=10, color="#F1F5F9")
        ax4.grid(True, axis="y")
        ax4.legend(loc="upper right", fontsize=8, facecolor="#1E293B", edgecolor="#334155")
    else:
        ax4.text(0.5, 0.5, "No Zone Data Available", ha="center", va="center", color="#94A3B8")

    # 5. WORKER LOCALIZATION & RSSI TRACKING
    ax5 = fig.add_subplot(gs[2, 0])
    if "RSSI (dBm)" in df.columns:
        valid_rssi_df = df[df["RSSI (dBm)"] > -900]
        if not valid_rssi_df.empty:
            ax5.plot(valid_rssi_df.index, valid_rssi_df["RSSI (dBm)"], color="#38BDF8", linewidth=1.4, label="Worker Wearable RSSI")
            ax5.axhline(-65, color="#22C55E", linestyle=":", linewidth=1.1, label="Strong (> -65 dBm / Near)")
            ax5.axhline(-80, color="#F59E0B", linestyle=":", linewidth=1.1, label="Moderate (-80 to -65 dBm)")
            ax5.axhline(-90, color="#EF4444", linestyle="--", linewidth=1.1, label="Weak / Dropout Danger")
            ax5.set_ylabel("Signal Strength (dBm)", fontsize=10)
            ax5.set_xlabel("Telemetry Cycle Index", fontsize=10)
        else:
            ax5.text(0.5, 0.5, "No Worker RSSI Detected", ha="center", va="center", color="#94A3B8")
    ax5.set_title("Worker Wearable Signal Health & Proximity", fontsize=13, fontweight="bold", pad=10, color="#F1F5F9")
    ax5.grid(True)
    ax5.legend(loc="lower left", fontsize=8, facecolor="#1E293B", edgecolor="#334155")

    # 6. HAZARD TRIGGER DISTRIBUTION
    ax6 = fig.add_subplot(gs[2, 1])
    if "Hazard Category" in df.columns:
        hazard_counts = df[df["Hazard Category"] != "Normal Air & Structure"]["Hazard Category"].value_counts().head(5)
        if not hazard_counts.empty:
            y_pos = np.arange(len(hazard_counts))
            bars = ax6.barh(y_pos, hazard_counts.values, color="#F43F5E", edgecolor="#9F1239", height=0.55)
            ax6.set_yticks(y_pos)
            ax6.set_yticklabels(hazard_counts.index, fontsize=9)
            ax6.invert_yaxis()
            ax6.set_xlabel("Incident Count", fontsize=10)
            for bar in bars:
                w = bar.get_width()
                ax6.text(w + max(hazard_counts.values)*0.02, bar.get_y() + bar.get_height()/2, f"{int(w)}",
                         va="center", ha="left", color="#F8FAFC", fontsize=9, fontweight="bold")
        else:
            ax6.text(0.5, 0.5, "No Hazard Events (All Normal)", ha="center", va="center", color="#22C55E", fontsize=12)
    ax6.set_title("Top Hazard Triggers & Threat Incidents", fontsize=13, fontweight="bold", pad=10, color="#F1F5F9")
    ax6.grid(True, axis="x")

    plt.savefig(output_path, dpi=200, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    print(f"[VISUAL GRAPHS] Successfully exported visual analytics dashboard to: {output_path}")

    # Also sync dashboard to legacy output path
    if sync_legacy:
        legacy_png = os.path.join(INSIGHTS_DIR, "mine_safety_dashboard.png")
        if output_path != legacy_png:
            try:
                import shutil
                shutil.copy2(output_path, legacy_png)
                print(f"[VISUAL GRAPHS] Synchronized visual dashboard to: {legacy_png}")
            except Exception as copy_err:
                print(f"[VISUAL GRAPHS] Note: Could not copy dashboard to legacy path: {copy_err}")

    return output_path


def draw_insights_to_excel(df_raw, target_path=EXCEL_FILE, sync_legacy=True):
    """
    Generates a professionally formatted, multi-tab Excel workbook
    drawing extensive analytics and rescue insights from mine telemetry data.
    """
    if df_raw is None or df_raw.empty:
        print("[EXCEL INSIGHTS] Warning: Dataframe is empty, skipping insight generation.")
        return False

    df = df_raw.copy()

    # Standardize column names and types
    for col in ["CH4 (ppm)", "CO (ppm)", "Vibration", "RSSI (dBm)", "CH4 Voltage (V)", "CO Voltage (V)"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    if "Zone" in df.columns:
        df["Zone"] = pd.to_numeric(df["Zone"], errors="coerce").fillna(0).astype(int)

    # Ensure Timestamp is fully populated
    if "Timestamp" not in df.columns or df["Timestamp"].isna().all():
        df["Timestamp"] = [
            (datetime.datetime.now() - datetime.timedelta(seconds=(len(df) - 1 - i) * 5)).strftime("%Y-%m-%d %H:%M:%S")
            for i in range(len(df))
        ]
    else:
        base_time = datetime.datetime.now()
        for idx in df.index:
            if pd.isna(df.at[idx, "Timestamp"]) or str(df.at[idx, "Timestamp"]).strip() in ("", "nan", "None"):
                df.at[idx, "Timestamp"] = (base_time - datetime.timedelta(seconds=(len(df) - 1 - idx) * 5)).strftime("%Y-%m-%d %H:%M:%S")

    # Compute Hazard Category and Safety Directive for all rows (including historical data)
    df["Hazard Category"] = [
        determine_hazard_category(
            float(r.get("CO (ppm)", 0) if pd.notna(r.get("CO (ppm)")) else 0),
            float(r.get("CH4 (ppm)", 0) if pd.notna(r.get("CH4 (ppm)")) else 0),
            float(r.get("Vibration", 0) if pd.notna(r.get("Vibration")) else 0),
            str(r.get("Risk Status", "SAFE") if pd.notna(r.get("Risk Status")) else "SAFE")
        )
        for _, r in df.iterrows()
    ]

    df["Safety Directive"] = [
        determine_safety_action(
            str(r.get("Risk Status", "SAFE") if pd.notna(r.get("Risk Status")) else "SAFE"),
            r.get("Hazard Category", "")
        )
        for _, r in df.iterrows()
    ]

    # --------------------------------------------------------
    # 1. GENERATE ANALYTICAL DATA STRUCTURES
    # --------------------------------------------------------
    total_records = len(df)
    safe_count = int((df["Risk Status"] == "SAFE").sum())
    warning_count = int((df["Risk Status"] == "WARNING").sum())
    critical_count = int((df["Risk Status"] == "CRITICAL").sum())
    safety_score = round((safe_count / total_records * 100), 2) if total_records > 0 else 100.0

    max_co = float(df["CO (ppm)"].max()) if "CO (ppm)" in df.columns and not df["CO (ppm)"].isna().all() else 0.0
    mean_co = float(df["CO (ppm)"].mean()) if "CO (ppm)" in df.columns and not df["CO (ppm)"].isna().all() else 0.0
    max_ch4 = float(df["CH4 (ppm)"].max()) if "CH4 (ppm)" in df.columns and not df["CH4 (ppm)"].isna().all() else 0.0
    mean_ch4 = float(df["CH4 (ppm)"].mean()) if "CH4 (ppm)" in df.columns and not df["CH4 (ppm)"].isna().all() else 0.0
    max_vib = float(df["Vibration"].max()) if "Vibration" in df.columns and not df["Vibration"].isna().all() else 0.0
    mean_vib = float(df["Vibration"].mean()) if "Vibration" in df.columns and not df["Vibration"].isna().all() else 0.0

    valid_rssi = df["RSSI (dBm)"].dropna()
    valid_rssi = valid_rssi[valid_rssi > -900]
    avg_rssi = float(valid_rssi.mean()) if not valid_rssi.empty else -999.0

    # Executive Summary Sheet Data
    exec_summary_kpis = [
        ["Total Telemetry Readings Logged", total_records, "Total telemetry cycles processed from wireless nodes"],
        ["Overall Mine Safety Health Score", f"{safety_score:.1f}%", "Percentage of operational time in SAFE condition"],
        ["Total SAFE Readings", safe_count, "Records within safe operational environmental limits"],
        ["Total WARNING Incidents", warning_count, "Records requiring proactive ventilation or monitoring"],
        ["Total CRITICAL Emergencies", critical_count, "Records demanding immediate worker alert & evacuation"],
        ["Peak Carbon Monoxide (CO)", f"{max_co:.2f} ppm", f"Threshold: Warn >={CO_WARNING} ppm, Crit >={CO_CRITICAL} ppm (Mean: {mean_co:.2f} ppm)"],
        ["Peak Methane Gas (CH4)", f"{max_ch4:.2f} ppm", f"Threshold: Warn >={CH4_WARNING} ppm, Crit >={CH4_CRITICAL} ppm (Mean: {mean_ch4:.2f} ppm)"],
        ["Peak Seismic Vibration", f"{max_vib:.4f} g", f"Threshold: Warn >={VIB_WARNING} g, Crit >={VIB_CRITICAL} g (Mean: {mean_vib:.4f} g)"],
        ["Average Worker Signal (RSSI)", f"{avg_rssi:.1f} dBm" if avg_rssi > -900 else "No Signal", "Worker wearable LoRa link signal quality indicator"],
        ["Report Generation Timestamp", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "Real-time AI Model analytical synchronization"]
    ]

    df_exec_kpi = pd.DataFrame(exec_summary_kpis, columns=["Operational Safety Metric", "Recorded Value", "Context / Engineering Benchmark"])

    # Key Dynamic Executive Findings & Automated Directives
    findings = []
    if critical_count > 0:
        findings.append(f"URGENT: {critical_count} critical hazard instances recorded. Active safety measures must be audited.")
    else:
        findings.append("Operational safety is stable with 0 critical emergencies currently logged.")

    if max_ch4 >= CH4_WARNING:
        findings.append(f"Methane levels surged to a maximum of {max_ch4:.2f} ppm. Flammable gas buildup requires booster fan engagement.")
    if max_co >= CO_WARNING:
        findings.append(f"Carbon Monoxide reached {max_co:.2f} ppm. Combustion / smoldering source inspection recommended.")
    if max_vib >= VIB_WARNING:
        findings.append(f"Seismic vibration peaked at {max_vib:.4f} g. Structural tunnel integrity checks advised along fault zones.")
    if avg_rssi > -65:
        findings.append("Worker wearable signal is STRONG, indicating close proximity to node receivers.")
    elif avg_rssi > -80:
        findings.append("Worker wearable signal is MODERATE across monitored tunnel sectors.")
    elif avg_rssi > -900:
        findings.append("Worker wearable signal is WEAK; worker is nearing peripheral tunnel boundary.")

    df_findings = pd.DataFrame({
        "Key Analytical Findings & AI Diagnostics": findings
    })

    # Recommended Rescue & Operational Actions
    rescue_actions = [
        ["Ventilation Control", "Maintain continuous air sweep; increase extraction if CH4 > 40 ppm or CO > 35 ppm."],
        ["Evacuation Readiness", "Keep emergency escape routes along Zone 1 clear of obstructions with lit pathway markers."],
        ["Wearable Watchdog", "Ensure worker panic button & bidirectional acoustic buzzer alerts remain synchronized."],
        ["Sensor Calibration", "Conduct baseline auto-zeroing routine on MQ-7 and MQ-4 gas sensors every shift."]
    ]
    df_actions = pd.DataFrame(rescue_actions, columns=["Rescue & Safety Protocol", "Standard Operating Procedure"])

    # --------------------------------------------------------
    # 2. ZONE RISK ANALYSIS SHEET
    # --------------------------------------------------------
    zone_metrics = []
    for zid in sorted(df["Zone"].unique()):
        if zid == 0:
            continue
        zdf = df[df["Zone"] == zid]
        z_total = len(zdf)
        z_safe = int((zdf["Risk Status"] == "SAFE").sum())
        z_warn = int((zdf["Risk Status"] == "WARNING").sum())
        z_crit = int((zdf["Risk Status"] == "CRITICAL").sum())

        z_safe_pct = round((z_safe / z_total * 100), 1) if z_total > 0 else 0
        z_warn_pct = round((z_warn / z_total * 100), 1) if z_total > 0 else 0
        z_crit_pct = round((z_crit / z_total * 100), 1) if z_total > 0 else 0

        co_m = float(zdf["CO (ppm)"].mean()) if not zdf["CO (ppm)"].isna().all() else 0.0
        co_mx = float(zdf["CO (ppm)"].max()) if not zdf["CO (ppm)"].isna().all() else 0.0
        co_std = float(zdf["CO (ppm)"].std()) if z_total > 1 else 0.0

        ch4_m = float(zdf["CH4 (ppm)"].mean()) if not zdf["CH4 (ppm)"].isna().all() else 0.0
        ch4_mx = float(zdf["CH4 (ppm)"].max()) if not zdf["CH4 (ppm)"].isna().all() else 0.0
        ch4_std = float(zdf["CH4 (ppm)"].std()) if z_total > 1 else 0.0

        vib_m = float(zdf["Vibration"].mean()) if not zdf["Vibration"].isna().all() else 0.0
        vib_mx = float(zdf["Vibration"].max()) if not zdf["Vibration"].isna().all() else 0.0

        z_rssi = zdf["RSSI (dBm)"].dropna()
        z_rssi = z_rssi[z_rssi > -900]
        rssi_m = float(z_rssi.mean()) if not z_rssi.empty else -999.0

        threat = "NORMAL"
        if z_crit > 0:
            threat = "HIGH HAZARD / CRITICAL"
        elif z_warn > 0:
            threat = "MODERATE / CAUTION"

        zone_metrics.append({
            "Mine Zone": f"Zone {zid}",
            "Readings": z_total,
            "SAFE Count": z_safe,
            "SAFE %": f"{z_safe_pct}%",
            "WARNING Count": z_warn,
            "WARNING %": f"{z_warn_pct}%",
            "CRITICAL Count": z_crit,
            "CRITICAL %": f"{z_crit_pct}%",
            "CO Mean (ppm)": round(co_m, 2),
            "CO Max (ppm)": round(co_mx, 2),
            "CO Std Dev": round(co_std, 2),
            "CH4 Mean (ppm)": round(ch4_m, 2),
            "CH4 Max (ppm)": round(ch4_mx, 2),
            "CH4 Std Dev": round(ch4_std, 2),
            "Vib Mean (g)": round(vib_m, 4),
            "Vib Max (g)": round(vib_mx, 4),
            "Mean RSSI (dBm)": round(rssi_m, 1) if rssi_m > -900 else "N/A",
            "Zone Hazard Evaluation": threat
        })

    df_zone_analysis = pd.DataFrame(zone_metrics)

    # --------------------------------------------------------
    # 3. CRITICAL HAZARD INCIDENTS LOG
    # --------------------------------------------------------
    df_incidents = df[
        (df["Risk Status"].isin(["WARNING", "CRITICAL"])) |
        (df["CO (ppm)"] >= CO_WARNING) |
        (df["CH4 (ppm)"] >= CH4_WARNING) |
        (df["Vibration"] >= VIB_WARNING)
    ].copy()

    incident_cols = [
        "Timestamp", "Zone", "Risk Status", "Hazard Category",
        "CO (ppm)", "CH4 (ppm)", "Vibration", "RSSI (dBm)", "Direction", "Safety Directive"
    ]
    incident_cols_present = [c for c in incident_cols if c in df_incidents.columns]
    df_incidents = df_incidents[incident_cols_present] if not df_incidents.empty else pd.DataFrame(columns=incident_cols)

    # --------------------------------------------------------
    # 4. WORKER PROXIMITY & EXPOSURE SHEET
    # --------------------------------------------------------
    proximity_records = []
    for zid in sorted(df["Zone"].unique()):
        if zid == 0:
            continue
        zdf = df[df["Zone"] == zid]
        valid_z_rssi = zdf["RSSI (dBm)"].dropna()
        valid_z_rssi = valid_z_rssi[valid_z_rssi > -900]

        strong_count = int((valid_z_rssi >= -65).sum())
        moderate_count = int(((valid_z_rssi < -65) & (valid_z_rssi >= -80)).sum())
        weak_count = int((valid_z_rssi < -80).sum())

        danger_exposure = int(((zdf["Risk Status"].isin(["WARNING", "CRITICAL"])) & (zdf["RSSI (dBm)"] >= -75)).sum())

        proximity_records.append({
            "Mine Zone": f"Zone {zid}",
            "Worker Detected Cycles": len(valid_z_rssi),
            "Close Proximity (RSSI >= -65 dBm)": strong_count,
            "Mid Range (-65 to -80 dBm)": moderate_count,
            "Far Range (< -80 dBm)": weak_count,
            "High-Hazard Proximity Cycles": danger_exposure,
            "Proximity Risk Index": "ELEVATED EXPOSURE" if danger_exposure > 0 else "SAFE DISTANCE"
        })

    df_proximity = pd.DataFrame(proximity_records)

    # --------------------------------------------------------
    # 5. GENERATE VISUAL DASHBOARD PNG GRAPHS FIRST
    # --------------------------------------------------------
    dashboard_png = os.path.join(INSIGHTS_DIR, "mine_rescue_dashboard.png")
    try:
        generate_visual_dashboard(df, output_path=dashboard_png, sync_legacy=sync_legacy)
    except Exception as dash_err:
        print(f"[VISUAL GRAPHS ERROR] Could not generate dashboard PNG: {dash_err}")

    # --------------------------------------------------------
    # 6. WRITE MULTI-SHEET EXCEL WORKBOOK & APPLY STYLING
    # --------------------------------------------------------
    active_target = target_path
    try:
        try:
            writer = pd.ExcelWriter(active_target, engine="openpyxl")
        except PermissionError:
            print(f"[EXCEL LOCK WARNING] '{target_path}' is currently open in Excel. Saving to alternative file...")
            active_target = target_path.replace(".xlsx", "_latest.xlsx")
            writer = pd.ExcelWriter(active_target, engine="openpyxl")

        with writer:
            # Sheet 1: Analytics Dashboard (Placeholder DataFrame, populated below)
            pd.DataFrame().to_excel(writer, sheet_name="Analytics Dashboard", index=False)
            # Sheet 2: Executive Summary
            df_exec_kpi.to_excel(writer, sheet_name="Executive Summary", index=False, startrow=1)
            # Sheet 3: Zone Risk Analysis
            df_zone_analysis.to_excel(writer, sheet_name="Zone Risk Analysis", index=False)
            # Sheet 4: Hazard Incidents
            df_incidents.to_excel(writer, sheet_name="Hazard Incidents Log", index=False)
            # Sheet 5: Worker Proximity
            df_proximity.to_excel(writer, sheet_name="Worker Proximity & Exposure", index=False)
            # Sheet 6: Raw Telemetry
            df.to_excel(writer, sheet_name="Sensor Data", index=False)

        # Apply OpenPyXL Visual Formatting
        wb = openpyxl.load_workbook(active_target)

        # Style Definitions
        font_main_title = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
        font_title = Font(name="Segoe UI", size=14, bold=True, color="1F4E79")
        font_sub_title = Font(name="Segoe UI", size=10, italic=True, color="E2E8F0")
        font_section = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        font_regular = Font(name="Segoe UI", size=10, color="000000")
        font_bold = Font(name="Segoe UI", size=10, bold=True, color="000000")
        font_kpi_val = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        font_kpi_lbl = Font(name="Segoe UI", size=9, bold=True, color="E2E8F0")

        fill_banner = PatternFill(start_color="0B1220", end_color="0B1220", fill_type="solid")
        fill_sub_banner = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        fill_header = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        fill_sub_header = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

        # KPI Card Fills
        fill_kpi_green = PatternFill(start_color="15803D", end_color="15803D", fill_type="solid")
        fill_kpi_blue = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
        fill_kpi_red = PatternFill(start_color="B91C1C", end_color="B91C1C", fill_type="solid")
        fill_kpi_amber = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
        fill_kpi_purple = PatternFill(start_color="6D28D9", end_color="6D28D9", fill_type="solid")

        fill_safe = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
        font_safe = Font(name="Segoe UI", size=10, bold=True, color="155724")

        fill_warn = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
        font_warn = Font(name="Segoe UI", size=10, bold=True, color="856404")

        fill_crit = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
        font_crit = Font(name="Segoe UI", size=10, bold=True, color="721C24")

        thin_side = Side(style="thin", color="D9D9D9")
        border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        # ----------------------------------------------------
        # 1. BUILD SHEET 1: ANALYTICS DASHBOARD
        # ----------------------------------------------------
        ws_dash = wb["Analytics Dashboard"]
        ws_dash.views.sheetView[0].showGridLines = True

        # Header Title
        ws_dash.merge_cells("A1:L1")
        cell_t1 = ws_dash["A1"]
        cell_t1.value = "⛏️ UNDERGROUND MINE SAFETY & RESCUE ANALYTICS DASHBOARD"
        cell_t1.font = font_main_title
        cell_t1.fill = fill_banner
        cell_t1.alignment = Alignment(horizontal="center", vertical="center")

        ws_dash.merge_cells("A2:L2")
        cell_t2 = ws_dash["A2"]
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cell_t2.value = f"Live IoT Telemetry, AI Hazard Prediction & Structural Diagnostics | Generated: {now_str}"
        cell_t2.font = font_sub_title
        cell_t2.fill = fill_sub_banner
        cell_t2.alignment = Alignment(horizontal="center", vertical="center")

        ws_dash.row_dimensions[1].height = 28
        ws_dash.row_dimensions[2].height = 20

        # KPI Metric Cards (Row 4: Title/Label, Row 5: Value, Row 6: Subtext)
        kpi_cards = [
            {"cols": ("A", "B"), "label": "MINE SAFETY RATING", "val": f"{safety_score:.1f}%", "sub": "Overall Health Rating", "fill": fill_kpi_green if safety_score > 60 else (fill_kpi_amber if safety_score > 40 else fill_kpi_red)},
            {"cols": ("C", "D"), "label": "TOTAL TELEMETRY CYCLES", "val": f"{total_records:,}", "sub": "Data Records Logged", "fill": fill_kpi_blue},
            {"cols": ("E", "F"), "label": "CRITICAL HAZARD EVENTS", "val": f"{critical_count:,}", "sub": f"{critical_count/total_records*100:.1f}% Critical Rate" if total_records else "0%", "fill": fill_kpi_red},
            {"cols": ("G", "H"), "label": "PEAK CARBON MONOXIDE", "val": f"{max_co:.1f} ppm", "sub": f"Warn >={CO_WARNING:.0f}, Crit >={CO_CRITICAL:.0f}", "fill": fill_kpi_amber if max_co >= CO_WARNING else fill_kpi_blue},
            {"cols": ("I", "J"), "label": "PEAK METHANE (CH4)", "val": f"{max_ch4:.1f} ppm", "sub": f"Warn >={CH4_WARNING:.0f}, Crit >={CH4_CRITICAL:.0f}", "fill": fill_kpi_amber if max_ch4 >= CH4_WARNING else fill_kpi_blue},
            {"cols": ("K", "L"), "label": "MAX SEISMIC VIBRATION", "val": f"{max_vib:.4f} g", "sub": f"Warn >={VIB_WARNING:.1f}g, Crit >={VIB_CRITICAL:.1f}g", "fill": fill_kpi_purple if max_vib >= VIB_WARNING else fill_kpi_blue},
        ]

        for card in kpi_cards:
            c1, c2 = card["cols"]
            ws_dash.merge_cells(f"{c1}4:{c2}4")
            lbl_c = ws_dash[f"{c1}4"]
            lbl_c.value = card["label"]
            lbl_c.font = font_kpi_lbl
            lbl_c.fill = card["fill"]
            lbl_c.alignment = Alignment(horizontal="center", vertical="center")

            ws_dash.merge_cells(f"{c1}5:{c2}5")
            val_c = ws_dash[f"{c1}5"]
            val_c.value = card["val"]
            val_c.font = font_kpi_val
            val_c.fill = card["fill"]
            val_c.alignment = Alignment(horizontal="center", vertical="center")

            ws_dash.merge_cells(f"{c1}6:{c2}6")
            sub_c = ws_dash[f"{c1}6"]
            sub_c.value = card["sub"]
            sub_c.font = Font(name="Segoe UI", size=8, color="E2E8F0")
            sub_c.fill = card["fill"]
            sub_c.alignment = Alignment(horizontal="center", vertical="center")

        ws_dash.row_dimensions[4].height = 18
        ws_dash.row_dimensions[5].height = 26
        ws_dash.row_dimensions[6].height = 16

        # Section Banner for Dashboard Image
        ws_dash.merge_cells("A8:L8")
        cell_img_title = ws_dash["A8"]
        cell_img_title.value = "📈 MULTI-PARAMETRIC VISUAL ANALYTICS & TELEMETRY PROGRESSION"
        cell_img_title.font = font_section
        cell_img_title.fill = fill_sub_header
        cell_img_title.alignment = Alignment(horizontal="center", vertical="center")
        ws_dash.row_dimensions[8].height = 22

        # Embed Dashboard PNG Image
        if os.path.exists(dashboard_png):
            try:
                img = OpenPyXLImage(dashboard_png)
                img.width = 980
                img.height = 650
                ws_dash.add_image(img, "A9")
            except Exception as img_err:
                print(f"[EXCEL IMAGE] Note: Could not insert image: {img_err}")

        # Section Banner for Cross-Zone Table below image
        row_table_start = 45
        ws_dash.merge_cells(f"A{row_table_start}:L{row_table_start}")
        cell_tbl_title = ws_dash[f"A{row_table_start}"]
        cell_tbl_title.value = "📋 CROSS-ZONE REAL-TIME TELEMETRY & HEALTH SCORECARD"
        cell_tbl_title.font = font_section
        cell_tbl_title.fill = fill_sub_header
        cell_tbl_title.alignment = Alignment(horizontal="center", vertical="center")
        ws_dash.row_dimensions[row_table_start].height = 22

        # Populate Zone Summary on Dashboard
        r_head = row_table_start + 1
        dash_zone_cols = ["Mine Zone", "Readings", "SAFE %", "WARNING %", "CRITICAL %", "CO Mean (ppm)", "CO Max (ppm)", "CH4 Mean (ppm)", "CH4 Max (ppm)", "Vib Max (g)", "Mean RSSI", "Hazard Evaluation"]
        for c_idx, col_name in enumerate(dash_zone_cols, 1):
            cell = ws_dash.cell(row=r_head, column=c_idx, value=col_name)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, z_row in df_zone_analysis.iterrows():
            row_num = r_head + 1 + r_idx
            row_data = [
                z_row.get("Mine Zone", ""),
                z_row.get("Readings", 0),
                z_row.get("SAFE %", "0%"),
                z_row.get("WARNING %", "0%"),
                z_row.get("CRITICAL %", "0%"),
                z_row.get("CO Mean (ppm)", 0),
                z_row.get("CO Max (ppm)", 0),
                z_row.get("CH4 Mean (ppm)", 0),
                z_row.get("CH4 Max (ppm)", 0),
                z_row.get("Vib Max (g)", 0),
                z_row.get("Mean RSSI (dBm)", "N/A"),
                z_row.get("Zone Hazard Evaluation", "NORMAL")
            ]
            for c_idx, val in enumerate(row_data, 1):
                cell = ws_dash.cell(row=row_num, column=c_idx, value=val)
                cell.font = font_regular
                cell.border = border_cell
                if r_idx % 2 == 1:
                    cell.fill = fill_zebra
                val_str = str(val).upper()
                if "CRITICAL" in val_str or "HIGH HAZARD" in val_str:
                    cell.fill = fill_crit
                    cell.font = font_crit
                elif "WARNING" in val_str or "CAUTION" in val_str:
                    cell.fill = fill_warn
                    cell.font = font_warn
                elif "SAFE" in val_str or "NORMAL" in val_str:
                    cell.fill = fill_safe
                    cell.font = font_safe

        # ----------------------------------------------------
        # 2. BUILD SHEET 2: EXECUTIVE SUMMARY
        # ----------------------------------------------------
        ws_exec = wb["Executive Summary"]
        ws_exec.views.sheetView[0].showGridLines = True
        ws_exec.cell(row=1, column=1, value="⛏️ MINE SAFETY & RESCUE EXECUTIVE REPORT").font = font_title

        # Style KPI Table
        for col_idx in range(1, len(df_exec_kpi.columns) + 1):
            cell = ws_exec.cell(row=2, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx in range(3, 3 + len(df_exec_kpi)):
            for col_idx in range(1, len(df_exec_kpi.columns) + 1):
                cell = ws_exec.cell(row=row_idx, column=col_idx)
                cell.font = font_bold if col_idx == 2 else font_regular
                cell.border = border_cell
                if row_idx % 2 == 0:
                    cell.fill = fill_zebra

        # Append Findings and Recommendations to Executive Summary
        start_f = 3 + len(df_exec_kpi) + 2
        ws_exec.cell(row=start_f, column=1, value="🔍 KEY AI DIAGNOSTIC FINDINGS").font = font_title
        start_f += 1
        ws_exec.cell(row=start_f, column=1, value="Finding Description").font = font_section
        ws_exec.cell(row=start_f, column=1).fill = fill_sub_header
        ws_exec.merge_cells(start_row=start_f, start_column=1, end_row=start_f, end_column=3)

        for i, item in enumerate(findings):
            r = start_f + 1 + i
            ws_exec.cell(row=r, column=1, value=f"• {item}").font = font_regular
            ws_exec.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            ws_exec.cell(row=r, column=1).border = border_cell

        start_a = start_f + len(findings) + 2
        ws_exec.cell(row=start_a, column=1, value="📋 ACTIONABLE RESCUE DIRECTIVES & SOPS").font = font_title
        start_a += 1
        for col_idx, h_text in enumerate(["Protocol Name", "Standard Operating Procedure Directive", "Status"]):
            cell = ws_exec.cell(row=start_a, column=col_idx + 1, value=h_text)
            cell.font = font_section
            cell.fill = fill_sub_header
            cell.border = border_cell

        for i, act in enumerate(rescue_actions):
            r = start_a + 1 + i
            c1 = ws_exec.cell(row=r, column=1, value=act[0])
            c2 = ws_exec.cell(row=r, column=2, value=act[1])
            c3 = ws_exec.cell(row=r, column=3, value="ACTIVE / ENFORCED")
            for c in [c1, c2, c3]:
                c.font = font_regular
                c.border = border_cell
                if r % 2 == 0:
                    c.fill = fill_zebra

        # Style all remaining sheets
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            ws.views.sheetView[0].showGridLines = True

            if sheet_name not in ["Executive Summary", "Analytics Dashboard"]:
                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.font = font_header
                    cell.fill = fill_header
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                for row_idx in range(2, ws.max_row + 1):
                    for col_idx in range(1, ws.max_column + 1):
                        cell = ws.cell(row=row_idx, column=col_idx)
                        cell.font = font_regular
                        cell.border = border_cell

                        if row_idx % 2 == 1:
                            cell.fill = fill_zebra

                        val_str = str(cell.value).strip().upper() if cell.value is not None else ""
                        if val_str == "CRITICAL" or "CRITICAL" in val_str or "HIGH HAZARD" in val_str:
                            cell.fill = fill_crit
                            cell.font = font_crit
                        elif val_str == "WARNING" or "WARNING" in val_str or "CAUTION" in val_str or "ELEVATED" in val_str:
                            cell.fill = fill_warn
                            cell.font = font_warn
                        elif val_str == "SAFE" or "NORMAL" in val_str:
                            cell.fill = fill_safe
                            cell.font = font_safe

            # Auto-adjust column width (skip Analytics Dashboard to keep card formatting clean)
            if sheet_name != "Analytics Dashboard":
                for col in ws.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        val = str(cell.value or "")
                        if "\n" in val:
                            val = max(val.split("\n"), key=len)
                        max_len = max(max_len, len(val))
                    ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 65)

        # ----------------------------------------------------
        # 3. EMBED NATIVE INTERACTIVE CHARTS
        # ----------------------------------------------------
        try:
            # Executive Summary: Risk Distribution Pie Chart
            pie = PieChart()
            labels = Reference(ws_exec, min_col=1, min_row=4, max_row=6)
            data_pie = Reference(ws_exec, min_col=2, min_row=3, max_row=6)
            pie.add_data(data_pie, titles_from_data=True)
            pie.set_categories(labels)
            pie.title = "Operational Risk Distribution"
            pie.width = 16
            pie.height = 10
            ws_exec.add_chart(pie, "E2")

            # Zone Risk Analysis: Clustered Bar Chart
            ws_zone = wb["Zone Risk Analysis"]
            if len(df_zone_analysis) > 0:
                bar_zone = BarChart()
                bar_zone.type = "col"
                bar_zone.style = 10
                bar_zone.title = "Zone Gas Concentrations (Mean & Peak)"
                bar_zone.y_axis.title = "PPM"
                bar_zone.x_axis.title = "Mine Zone"
                data_z = Reference(ws_zone, min_col=9, min_row=1, max_col=13, max_row=len(df_zone_analysis) + 1)
                cats_z = Reference(ws_zone, min_col=1, min_row=2, max_row=len(df_zone_analysis) + 1)
                bar_zone.add_data(data_z, titles_from_data=True)
                bar_zone.set_categories(cats_z)
                bar_zone.width = 16
                bar_zone.height = 10
                ws_zone.add_chart(bar_zone, "T2")

            # Worker Proximity & Exposure: Proximity Breakdown
            ws_prox = wb["Worker Proximity & Exposure"]
            if len(df_proximity) > 0:
                bar_prox = BarChart()
                bar_prox.type = "col"
                bar_prox.style = 11
                bar_prox.title = "Worker Proximity Range Breakdown"
                bar_prox.y_axis.title = "Telemetry Cycles"
                bar_prox.x_axis.title = "Mine Zone"
                data_p = Reference(ws_prox, min_col=3, min_row=1, max_col=6, max_row=len(df_proximity) + 1)
                cats_p = Reference(ws_prox, min_col=1, min_row=2, max_row=len(df_proximity) + 1)
                bar_prox.add_data(data_p, titles_from_data=True)
                bar_prox.set_categories(cats_p)
                bar_prox.width = 16
                bar_prox.height = 10
                ws_prox.add_chart(bar_prox, "I2")
        except Exception as chart_err:
            print(f"[EXCEL CHARTS] Note: Chart embedding skipped: {chart_err}")

        # Save primary target
        try:
            wb.save(active_target)
            print(f"[EXCEL INSIGHTS] Successfully written comprehensive analytics to: {active_target}")
        except PermissionError:
            print(f"[EXCEL LOCK WARNING] '{active_target}' is currently open in Excel. Please close it to overwrite.")
            alt_path = active_target.replace(".xlsx", "_latest.xlsx")
            try:
                wb.save(alt_path)
                print(f"[EXCEL INSIGHTS] Saved insights to alternative file: {alt_path}")
            except Exception as alt_err:
                print(f"[EXCEL ERROR] Could not save alternative file: {alt_err}")

        # Also sync to mine_safety_data.xlsx if target is mine_rescue.xlsx
        if sync_legacy and target_path != LEGACY_EXCEL_FILE and os.path.exists(INSIGHTS_DIR):
            try:
                wb.save(LEGACY_EXCEL_FILE)
                print(f"[EXCEL INSIGHTS] Synchronized insights to legacy log: {LEGACY_EXCEL_FILE}")
            except PermissionError:
                print(f"[EXCEL LOCK WARNING] Legacy file '{LEGACY_EXCEL_FILE}' is open in Excel. Sync skipped.")
            except Exception as legacy_err:
                print(f"[EXCEL INSIGHTS] Note: Could not sync to legacy file: {legacy_err}")

        return True

    except Exception as e:
        print(f"[EXCEL INSIGHTS ERROR] Failed to write insights workbook: {e}")
        return False


# ============================================================
# INITIALIZE OR LOAD EXCEL LOG
# ============================================================

excel_data = None

# Attempt to load existing log from mine_rescue.xlsx or legacy mine_safety_data.xlsx
if os.path.exists(EXCEL_FILE):
    try:
        excel_data = pd.read_excel(EXCEL_FILE, sheet_name="Sensor Data")
        print(f"[EXCEL] Loaded existing telemetry from {EXCEL_FILE} ({len(excel_data)} rows)")
    except Exception:
        try:
            excel_data = pd.read_excel(EXCEL_FILE)
            print(f"[EXCEL] Loaded telemetry from {EXCEL_FILE} ({len(excel_data)} rows)")
        except Exception:
            excel_data = None

if excel_data is None and os.path.exists(LEGACY_EXCEL_FILE):
    try:
        excel_data = pd.read_excel(LEGACY_EXCEL_FILE, sheet_name="Sensor Data")
        print(f"[EXCEL] Migrated {len(excel_data)} historical rows from {LEGACY_EXCEL_FILE}")
    except Exception:
        try:
            excel_data = pd.read_excel(LEGACY_EXCEL_FILE)
            print(f"[EXCEL] Migrated {len(excel_data)} historical rows from {LEGACY_EXCEL_FILE}")
        except Exception:
            excel_data = None

if excel_data is None:
    excel_data = pd.DataFrame(columns=[
        "Timestamp",
        "Zone",
        "CH4 Voltage (V)",
        "CH4 (ppm)",
        "CO Voltage (V)",
        "CO (ppm)",
        "Vibration",
        "RSSI (dBm)",
        "Direction",
        "Risk Status",
        "ML Safe Prob",
        "ML Warning Prob",
        "ML Critical Prob",
        "Hazard Category",
        "Safety Directive"
    ])

# Ensure all columns exist
required_columns = [
    "Timestamp", "Zone", "CH4 Voltage (V)", "CH4 (ppm)", "CO Voltage (V)", "CO (ppm)",
    "Vibration", "RSSI (dBm)", "Direction", "Risk Status",
    "ML Safe Prob", "ML Warning Prob", "ML Critical Prob", "Hazard Category", "Safety Directive"
]

for col in required_columns:
    if col not in excel_data.columns:
        excel_data[col] = ""


# ============================================================
# PREVIOUS MPU DATA & VIBRATION
# ============================================================

previous_mpu = {1: None, 2: None}

def vibration_from_acceleration(zone_id, ax, ay, az):
    previous = previous_mpu[zone_id]
    if previous is None:
        previous_mpu[zone_id] = (ax, ay, az)
        return 0.0

    previous_ax, previous_ay, previous_az = previous
    delta_ax = ax - previous_ax
    delta_ay = ay - previous_ay
    delta_az = az - previous_az

    vibration = math.sqrt(delta_ax ** 2 + delta_ay ** 2 + delta_az ** 2)
    previous_mpu[zone_id] = (ax, ay, az)
    return vibration


# ============================================================
# PROCESS SENSOR SAMPLE
# ============================================================

def process_sensor_sample(zone_id, mq7_mV, mq4_mV, ax, ay, az, worker_rssi, direction="UNKNOWN"):
    global excel_data

    co_voltage = mq7_mV / 1000.0
    ch4_voltage = mq4_mV / 1000.0
    current_vibration = vibration_from_acceleration(zone_id, ax, ay, az)

    if not (0 <= co_voltage <= 3.3 and 0 <= ch4_voltage <= 3.3):
        print(f"ZONE {zone_id}: INVALID SENSOR VOLTAGE")
        return None

    current_co = float(mq7_voltage_to_co(co_voltage))
    current_ch4 = float(mq5_voltage_to_ch4(ch4_voltage))

    co_features = create_test_features(co_reference_gas, co_reference_vibration, current_co, current_vibration)
    ch4_features = create_test_features(ch4_reference_gas, ch4_reference_vibration, current_ch4, current_vibration)

    co_probabilities = get_probability(co_model, co_scaler, co_features)
    ch4_probabilities = get_probability(ch4_model, ch4_scaler, ch4_features)

    prediction, final_probabilities = fuse_predictions(co_probabilities, ch4_probabilities)

    hazard_category = determine_hazard_category(current_co, current_ch4, current_vibration, prediction)
    safety_directive = determine_safety_action(prediction, hazard_category)
    now_ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Append to Excel records
    new_row = pd.DataFrame([{
        "Timestamp": now_ts,
        "Zone": zone_id,
        "CH4 Voltage (V)": round(ch4_voltage, 3),
        "CH4 (ppm)": round(current_ch4, 2),
        "CO Voltage (V)": round(co_voltage, 3),
        "CO (ppm)": round(current_co, 2),
        "Vibration": round(current_vibration, 4),
        "RSSI (dBm)": worker_rssi,
        "Direction": direction,
        "Risk Status": prediction,
        "ML Safe Prob": round(final_probabilities.get("SAFE", 0.0), 3),
        "ML Warning Prob": round(final_probabilities.get("WARNING", 0.0), 3),
        "ML Critical Prob": round(final_probabilities.get("CRITICAL", 0.0), 3),
        "Hazard Category": hazard_category,
        "Safety Directive": safety_directive
    }])

    excel_data = pd.concat([excel_data, new_row], ignore_index=True)

    # Save and regenerate insights
    saved = draw_insights_to_excel(excel_data, EXCEL_FILE)
    excel_status = "SAVED (Insights Updated)" if saved else "ERROR"

    return {
        "zone": zone_id,
        "timestamp": now_ts,
        "ch4_voltage": ch4_voltage,
        "ch4": current_ch4,
        "co_voltage": co_voltage,
        "co": current_co,
        "vibration": current_vibration,
        "rssi": worker_rssi,
        "direction": direction,
        "risk": prediction,
        "probabilities": final_probabilities,
        "hazard_category": hazard_category,
        "safety_directive": safety_directive,
        "excel": excel_status
    }


# ============================================================
# ZONE AND WORKER LOGIC
# ============================================================

latest_zone_data = {1: None, 2: None}

def determine_approaching_zone():
    zone1 = latest_zone_data[1]
    zone2 = latest_zone_data[2]

    rssi1 = zone1["rssi"] if zone1 is not None else None
    rssi2 = zone2["rssi"] if zone2 is not None else None

    valid1 = rssi1 is not None and rssi1 > -900
    valid2 = rssi2 is not None and rssi2 > -900

    if not valid1 and not valid2:
        return None
    if valid1 and not valid2:
        return 1
    if valid2 and not valid1:
        return 2

    return 1 if rssi1 > rssi2 else 2


def determine_direction():
    approaching_zone = determine_approaching_zone()
    if approaching_zone is None:
        return None, "UNKNOWN"

    zone = latest_zone_data[approaching_zone]
    risk = zone.get("risk", "UNKNOWN") if zone else "UNKNOWN"

    if risk == "CRITICAL":
        direction = "CRITICAL"
    elif risk == "WARNING":
        direction = "WARNING"
    elif risk == "SAFE":
        direction = "SAFE"
    else:
        direction = "UNKNOWN"

    return approaching_zone, direction


def determine_buzzer(risk):
    if risk == "CRITICAL":
        return "HIGH"
    elif risk == "WARNING":
        return "SOFT"
    return "OFF"


def send_alert(ser, approaching_zone, risk, direction):
    if approaching_zone is None:
        return

    buzzer = determine_buzzer(risk)
    command = f"ALERT,ZONE={approaching_zone},RISK={risk},DIRECTION={direction},BUZZER={buzzer}\n"

    try:
        ser.write(command.encode("utf-8"))
        ser.flush()
        print()
        print(">>> SENT TO ZONE 2 GATEWAY:")
        print(command.strip())
        print(f"    Worker approaching: ZONE {approaching_zone}")
        print(f"    Direction: {direction}")
        print(f"    Buzzer: {buzzer}")
    except serial.SerialException as error:
        print("FAILED TO SEND ALERT:", error)


# ============================================================
# SERIAL PACKET PARSER
# ============================================================

def parse_zone_packet(line):
    line = line.strip()
    if line.startswith("PYTHON_DATA:"):
        line = line[len("PYTHON_DATA:"):]
    else:
        return None

    fields = {}
    for part in line.split(","):
        part = part.strip()
        if "=" in part:
            key, value = part.split("=", 1)
            fields[key.strip().upper()] = value.strip()

    if "ZONE" not in fields:
        return None

    try:
        zone_id = int(fields["ZONE"])
    except ValueError:
        return None

    if zone_id not in (1, 2):
        return None

    required = ["MQ7", "MQ4", "AX", "AY", "AZ", "WRSSI"]
    for key in required:
        if key not in fields:
            print(f"ZONE {zone_id}: MISSING {key}")
            return None

    try:
        mq7_mV = float(fields["MQ7"])
        mq4_mV = float(fields["MQ4"])
        ax = float(fields["AX"])
        ay = float(fields["AY"])
        az = float(fields["AZ"])
        worker_rssi = float(fields["WRSSI"])
    except ValueError:
        print(f"ZONE {zone_id}: INVALID NUMERIC DATA")
        return None

    try:
        avg_rssi = float(fields.get("AVG_RSSI", worker_rssi))
    except ValueError:
        avg_rssi = worker_rssi

    try:
        rssi_change = float(fields.get("RSSI_CHANGE", 0.0))
    except ValueError:
        rssi_change = 0.0

    direction = fields.get("DIR", "UNKNOWN").upper()

    return {
        "zone": zone_id,
        "mq7_mV": mq7_mV,
        "mq4_mV": mq4_mV,
        "ax": ax,
        "ay": ay,
        "az": az,
        "rssi": worker_rssi,
        "avg_rssi": avg_rssi,
        "rssi_change": rssi_change,
        "direction": direction
    }


# ============================================================
# PRINT CONSOLE INSIGHTS SUMMARY
# ============================================================

def print_insights_summary():
    global excel_data
    if excel_data is None or excel_data.empty:
        print("No telemetry data available for insight summary.")
        return

    print()
    print("=" * 70)
    print("           [REPORT] MINE RESCUE AI - COMPREHENSIVE INSIGHTS SUMMARY")
    print("=" * 70)

    total = len(excel_data)
    safe = int((excel_data["Risk Status"] == "SAFE").sum())
    warn = int((excel_data["Risk Status"] == "WARNING").sum())
    crit = int((excel_data["Risk Status"] == "CRITICAL").sum())
    score = (safe / total * 100) if total > 0 else 100

    print(f"Total Telemetry Cycles Processed : {total}")
    print(f"Mine Operational Safety Rating    : {score:.1f}%")
    print(f"  SAFE Cycles                     : {safe} ({safe/total*100:.1f}%)")
    print(f"  WARNING Incidents               : {warn} ({warn/total*100:.1f}%)")
    print(f"  CRITICAL Hazards                : {crit} ({crit/total*100:.1f}%)")
    print("-" * 70)

    co_vals = pd.to_numeric(excel_data["CO (ppm)"], errors="coerce")
    ch4_vals = pd.to_numeric(excel_data["CH4 (ppm)"], errors="coerce")
    vib_vals = pd.to_numeric(excel_data["Vibration"], errors="coerce")

    print(f"Carbon Monoxide (CO) Peak : {co_vals.max():.2f} ppm  (Mean: {co_vals.mean():.2f} ppm)")
    print(f"Methane (CH4) Peak        : {ch4_vals.max():.2f} ppm  (Mean: {ch4_vals.mean():.2f} ppm)")
    print(f"Seismic Vibration Peak    : {vib_vals.max():.4f} g   (Mean: {vib_vals.mean():.4f} g)")
    print("-" * 70)
    print(f"Excel Analytics Workbook  : {EXCEL_FILE}")
    if os.path.exists(LEGACY_EXCEL_FILE):
        print(f"Legacy Synchronized Log   : {LEGACY_EXCEL_FILE}")
    print("=" * 70)
    print()


# ============================================================
# MAIN ENTRYPOINT / CLI
# ============================================================

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Mine Safety Predictive AI & Rescue Analytics Engine")
    parser.add_argument("--insights", action="store_true", help="Generate multi-tab Excel insights report from existing logs and exit")
    parser.add_argument("--excel", type=str, default=EXCEL_FILE, help="Path to target Excel report file")
    parser.add_argument("--port", type=str, default=SERIAL_PORT, help="Serial COM port for ESP32 Gateway")
    args, unknown = parser.parse_known_args()

    if args.excel:
        EXCEL_FILE = args.excel

    # Check if user invoked offline insight generation
    if args.insights:
        print()
        print("=" * 60)
        print("GENERATING COMPREHENSIVE MINE RESCUE EXCEL INSIGHTS")
        print("=" * 60)
        draw_insights_to_excel(excel_data, EXCEL_FILE)
        print_insights_summary()
        sys.exit(0)

    # If telemetry log exists and we want an immediate sync before serial loop
    if not excel_data.empty:
        draw_insights_to_excel(excel_data, EXCEL_FILE)

    # --------------------------------------------------------
    # SERIAL HARDWARE CONNECTION
    # --------------------------------------------------------
    print()
    print("=" * 60)
    print("MINE SAFETY AI SYSTEM - LIVE SENSOR BRIDGE")
    print("=" * 60)
    print(f"Gateway port : {args.port}")
    print(f"Baud rate    : {BAUD_RATE}")
    print(f"Insights File: {EXCEL_FILE}")
    print()
    print("Close Arduino Serial Monitor before starting Python.")
    print()

    try:
        esp32 = serial.Serial(args.port, BAUD_RATE, timeout=1)
        time.sleep(2)
        print("CONNECTED TO ZONE 2 / MAIN GATEWAY")
        print("Waiting for sensor data...")
        print()
    except serial.SerialException as error:
        print(f"Could not open {args.port}")
        print("Make sure Arduino Serial Monitor is CLOSED.")
        print(f"Serial error: {error}")
        print()
        print("Hardware gateway offline. Generating comprehensive insights from logged telemetry data...")
        draw_insights_to_excel(excel_data, EXCEL_FILE)
        print_insights_summary()
        sys.exit(0)

    # --------------------------------------------------------
    # MAIN ML LIVE LOOP
    # --------------------------------------------------------
    next_ml_time = time.monotonic() + ML_INTERVAL_SECONDS

    try:
        while True:
            line = esp32.readline().decode("utf-8", errors="ignore").strip()
            if line:
                parsed = parse_zone_packet(line)
                if parsed is not None:
                    zone_id = parsed["zone"]
                    latest_zone_data[zone_id] = parsed

                    print()
                    print("[SERIAL RECEIVED]")
                    print(f"ZONE={zone_id}")
                    print(f"MQ7={parsed['mq7_mV']:.0f} mV")
                    print(f"MQ4={parsed['mq4_mV']:.0f} mV")
                    print(f"AX={parsed['ax']:.4f}, AY={parsed['ay']:.4f}, AZ={parsed['az']:.4f}")
                    print(f"WRSSI={parsed['rssi']:.0f} dBm, DIR={parsed['direction']}")
                    print(f"ZONE {zone_id} DATA STORED")

            current_time = time.monotonic()
            if current_time >= next_ml_time:
                print()
                print("=" * 70)
                print("                 ML UPDATE & INSIGHTS SYNC")
                print("=" * 70)

                for zone_id in (1, 2):
                    data = latest_zone_data[zone_id]
                    if data is None:
                        print(f"ZONE {zone_id} STATUS: NO DATA RECEIVED")
                        continue

                    result = process_sensor_sample(
                        zone_id,
                        data["mq7_mV"],
                        data["mq4_mV"],
                        data["ax"],
                        data["ay"],
                        data["az"],
                        data["rssi"],
                        data["direction"]
                    )

                    if result is None:
                        continue

                    latest_zone_data[zone_id]["risk"] = result["risk"]
                    probs = result["probabilities"]

                    print("-" * 42)
                    print(f"ZONE {zone_id}")
                    print(f"CO   : {result['co']:.2f} ppm")
                    print(f"CH4  : {result['ch4']:.2f} ppm")
                    print(f"VIB  : {result['vibration']:.4f} g")
                    print(f"RSSI : {result['rssi']:.0f} dBm")
                    print(f"RISK : {result['risk']} ({result['hazard_category']})")
                    print(f"ML PROBS -> SAFE: {probs.get('SAFE', 0):.3f}, WARN: {probs.get('WARNING', 0):.3f}, CRIT: {probs.get('CRITICAL', 0):.3f}")
                    print(f"Directive : {result['safety_directive']}")
                    print(f"Excel     : {result['excel']}")

                # Worker Direction and Alerting
                approaching_zone, direction = determine_direction()
                print()
                print("=" * 50)
                print("          WORKER SAFETY & RESCUE DECISION")
                print("=" * 50)

                if approaching_zone is None:
                    print("APPROACHING ZONE : UNKNOWN")
                    print("DIRECTION        : UNKNOWN")
                    print("BUZZER           : OFF")
                    command = "ALERT,ZONE=0,RISK=SAFE,DIRECTION=UNKNOWN,BUZZER=OFF\n"
                    try:
                        esp32.write(command.encode("utf-8"))
                        esp32.flush()
                        print(">>> SENT TO ZONE 2 GATEWAY: " + command.strip())
                    except serial.SerialException as error:
                        print("Alert send error:", error)
                else:
                    approaching_data = latest_zone_data[approaching_zone]
                    approaching_risk = approaching_data.get("risk", "UNKNOWN") if approaching_data else "UNKNOWN"
                    buzzer = determine_buzzer(approaching_risk)

                    print(f"APPROACHING ZONE : ZONE {approaching_zone}")
                    print(f"ZONE RISK        : {approaching_risk}")
                    print(f"DIRECTION        : {direction}")
                    print(f"BUZZER           : {buzzer}")

                    send_alert(esp32, approaching_zone, approaching_risk, direction)

                print("=" * 50)
                while next_ml_time <= current_time:
                    next_ml_time += ML_INTERVAL_SECONDS

    except KeyboardInterrupt:
        print()
        print("Stopping ML system...")
    finally:
        if 'esp32' in locals() and esp32.is_open:
            esp32.close()
            print("Serial connection closed.")
        print_insights_summary()
        print("ML system stopped.")